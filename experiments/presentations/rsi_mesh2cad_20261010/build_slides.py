"""Build a six-slide, editable PPTX plus matching PDF/PNG preview from saved results.
Run with the project Python and python-pptx; RSI_SLIDE_FONT points to a CJK font.
"""
import os, sys, json, math
from pathlib import Path
from decimal import Decimal
sys.path.insert(0, '/tmp/rsi-slide-deps')
ROOT=Path(__file__).resolve().parent
EXP=ROOT.parents[1]/'mesh_to_implicit_primitives'
sys.path.insert(0,str(EXP))
from PIL import Image, ImageDraw, ImageFont
import numpy as np
import trimesh
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from render_rsi_results import render

FONT=os.environ.get('RSI_SLIDE_FONT','/tmp/rsi-NotoSansCJKsc-Regular.otf')
W,H,S=1280,720,2
BG='#F5F7FB'; NAVY='#12243B'; BLUE='#176C9B'; TEAL='#007F78'; ORANGE='#B65424'; GRAY='#536477'; LIGHT='#E6EDF4'; WHITE='#FFFFFF'
prs=Presentation(); prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
prs.core_properties.title='基于 RSI 的 Mesh2CAD：阶段性汇报'
prs.core_properties.subject='问题定义、相关工作、方法流程、实验结果与未来计划'
prs.core_properties.author='Mesh2CAD 项目'
assets=ROOT/'assets'; assets.mkdir(exist_ok=True)
slides=[]; overflow=[]
refs={
 '1':('UCSG-Net — Unsupervised Discovering of Constructive Solid Geometry Tree','Kania, Zięba & Kajdanowicz, NeurIPS 2020','https://arxiv.org/abs/2006.09102','10.48550/arXiv.2006.09102'),
 '2':('CAD-Recode: Reverse Engineering CAD Code from Point Clouds','Rukhovich et al., ICCV 2025','https://arxiv.org/abs/2412.14042','10.48550/arXiv.2412.14042'),
 '3':('StepCAD: Mesh-to-CAD Code Generation via LLM Policy and Geometry-Guided Search','Nehme & Ahmed, arXiv 2026','https://arxiv.org/abs/2610.03799','10.48550/arXiv.2610.03799'),
 '4':('RSIAgent: Autonomous Exploration for Recursive Self-improvement in New Environments','Zhu et al., arXiv 2026','https://arxiv.org/abs/2609.15364','10.48550/arXiv.2609.15364')}

def rgb(c):return RGBColor.from_string(c.lstrip('#'))
def emu(v):return Inches(v/96)
def wrap(text,size,width):
 font=ImageFont.truetype(FONT,round(size*S));result=[]
 for paragraph in text.split('\n'):
  line=''
  for ch in paragraph:
   if line and font.getlength(line+ch)>width*S:
    result.append(line);line=ch
   else:line+=ch
  result.append(line)
 return result

class Slide:
 def __init__(self,num,section,title,subtitle=''):
  self.slide=prs.slides.add_slide(prs.slide_layouts[6]); self.im=Image.new('RGB',(W*S,H*S),BG);self.d=ImageDraw.Draw(self.im);self.num=num
  fill=self.slide.background.fill;fill.solid();fill.fore_color.rgb=rgb(BG)
  self.rect(0,0,12,H,BLUE)
  self.text(45,25,1150,24,f'RSI × MESH2CAD   /   {section}',16,TEAL)
  self.text(45,66,1190,60,title,36,NAVY,True)
  if subtitle:self.text(46,126,1180,52,subtitle,20,GRAY)
  self.text(45,686,1120,20,'阶段性研究汇报  ·  2026.10.10  ·  数值来自本地实验，非标准榜单成绩',13,GRAY)
  self.text(1180,684,60,24,f'{num:02d} / 06',15,BLUE)
  slides.append(self)
 def rect(self,x,y,w,h,color=WHITE,line=None,radius=False):
  self.d.rounded_rectangle([x*S,y*S,(x+w)*S,(y+h)*S],radius=10*S if radius else 0,fill=color,outline=line,width=S)
  sh=self.slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,emu(x),emu(y),emu(w),emu(h))
  sh.fill.solid();sh.fill.fore_color.rgb=rgb(color)
  if line:sh.line.color.rgb=rgb(line)
  else:sh.line.fill.background()
  return sh
 def text(self,x,y,w,h,text,size=22,color=NAVY,bold=False,link=None):
  lines=wrap(text,size,w); leading=size*1.35
  if len(lines)*leading>h+3:overflow.append((self.num,text[:40],len(lines)*leading,h))
  font=ImageFont.truetype(FONT,round(size*S))
  for i,line in enumerate(lines):self.d.text((x*S,(y+i*leading)*S),line,font=font,fill=color,anchor='lt',stroke_width=0)
  shape=self.slide.shapes.add_textbox(emu(x),emu(y),emu(w),emu(h)); tf=shape.text_frame;tf.word_wrap=False
  tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
  for i,line in enumerate(lines):
   p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.space_before=Pt(0);p.space_after=Pt(0);p.line_spacing=Pt(leading*.75)
   run=p.add_run();run.text=line;run.font.name='Microsoft YaHei';run.font.size=Pt(size*.75);run.font.bold=bold;run.font.color.rgb=rgb(color)
   ea=OxmlElement('a:ea');ea.set('typeface','Microsoft YaHei');run._r.get_or_add_rPr().append(ea)
   if link:run.hyperlink.address=link
  return shape
 def image(self,path,x,y,w,h):
  im=Image.open(path).convert('RGB'); scale=min(w/im.width,h/im.height); ww,hh=im.width*scale,im.height*scale
  xx,yy=x+(w-ww)/2,y+(h-hh)/2
  self.im.paste(im.resize((round(ww*S),round(hh*S)),Image.Resampling.LANCZOS),(round(xx*S),round(yy*S)))
  self.slide.shapes.add_picture(str(path),emu(xx),emu(yy),width=emu(ww),height=emu(hh))
 def note(self,text):self.slide.notes_slide.notes_text_frame.text=text
 def banner(self,y,text,color=TEAL):
  self.rect(45,y,1190,47,color,radius=True); self.text(62,y+9,1160,34,text,20,WHITE)

runs=EXP/'rsi_runs'; batch=runs/'standard_9_20261009_trial01';old=runs/'voron_motor_mount_trial01';new=runs/'voron_motor_mount_48_trial01'
def read_result(run,name='Voron_X_Motor_Mount_A'):return json.loads((run/'cases'/name/'result.json').read_text())
def cost(run):
 rows=[json.loads(s) for s in (run/'usage.jsonl').read_text().splitlines() if s]; value=Decimal(0);tokens=0
 for row in rows:
  u=row['usage'];d=u.get('prompt_tokens_details') or {};a=u['prompt_tokens'];b=u['completion_tokens'];c=d.get('cached_tokens') or 0;w=d.get('cache_write_tokens') or 0
  value+=(Decimal(a-c-w)*4+Decimal(c)*Decimal('.4')+Decimal(w)*5+Decimal(b)*20)/1000000;tokens+=u['total_tokens']
 return float(value),tokens
cohort=json.loads((batch/'cohort.json').read_text());results=[(x['name'],read_result(batch,x['name'])) for x in cohort]
a,b=read_result(old),read_result(new); ca,ta=cost(old);cb,tb=cost(new);cc,tc=cost(batch)
assert sum(v['status']=='completed' for _,v in results)==7

# Actual full triangle-mesh renderings; no generated imagery, no surface point clouds.
target=trimesh.load(EXP/'data/meshes/Voron_X_Motor_Mount_A.stl',force='mesh')
meshes=[target,trimesh.load(old/'cases/Voron_X_Motor_Mount_A/final/reconstruction.stl',force='mesh'),trimesh.load(new/'cases/Voron_X_Motor_Mount_A/final/reconstruction.stl',force='mesh')]
bounds=np.stack([m.bounds for m in meshes]);lo=bounds[:,0].min(0);hi=bounds[:,1].max(0);center=(lo+hi)/2;span=np.linalg.norm(hi-lo)*.55
rendered=[Image.fromarray(np.uint8(np.clip(render(m,center,span,25,-55,pixels=800),0,1)*255)) for m in meshes]
mask=np.any(np.stack([np.asarray(im) for im in rendered])<245,axis=(0,3))
yy,xx=np.where(mask)
box=(max(0,int(xx.min())-25),max(0,int(yy.min())-25),min(800,int(xx.max())+26),min(800,int(yy.max())+26))
for tag,im in zip(['voron_input','voron_24','voron_48'],rendered):
 im.crop(box).save(assets/f'{tag}.png')
hard='hard_109857_1d24326e_6'; original=trimesh.load(EXP/'fusion360_gallery/hard/109857_1d24326e_6.obj',force='mesh');recon=trimesh.load(batch/f'cases/{hard}/final/reconstruction.stl',force='mesh')
bounds=np.stack([original.bounds,recon.bounds]);lo=bounds[:,0].min(0);hi=bounds[:,1].max(0)
for tag,m in [('hard_input',original),('hard_output',recon)]:
 p=assets/f'{tag}.png'
 if not p.exists():Image.fromarray(np.uint8(np.clip(render(m,(lo+hi)/2,np.linalg.norm(hi-lo)*.55,25,-55,pixels=600),0,1)*255)).crop((100,100,510,510)).save(p)

s=Slide(1,'问题定义','从三角网格，到可编辑的几何程序','基于 RSI 的 Mesh2CAD：用执行反馈与经验记忆，搜索几何构造过程')
s.rect(45,195,560,395,WHITE,radius=True);s.text(68,215,510,40,'输入：STL / OBJ 三角网格',26,BLUE,True);s.image(assets/'voron_input.png',100,258,430,290)
s.text(70,551,510,30,'示例：Voron 电机支架的真实 mesh',18,GRAY)
s.text(632,308,45,50,'→',38,TEAL)
s.rect(700,195,535,395,WHITE,radius=True);s.text(725,217,485,40,'输出：参数化 CSG 树 + STL',26,BLUE,True)
s.text(728,278,472,170,'选择基元与布尔关系\n拟合位置、旋转和尺寸\n执行、验证、修改几何程序',25)
s.rect(725,439,485,121,LIGHT,radius=True);s.text(740,453,457,101,'六类基元 × 并 / 差 / 交\n目标：几何准确、结构简洁、结果可复现\n当前输出不含 STEP / 原始 CAD 时间线',19)
s.banner(610,'核心问题：固定模型参数，能否通过课程练习和记忆改善 mesh → CSG 重建？')
s.note('定位：当前实现是 mesh 到隐式 CSG 的阶段性原型，并非完整工业 CAD 历史恢复。问题是研究问题，不是已证明结论。输入示例为 mesh_to_implicit_primitives/data/meshes/Voron_X_Motor_Mount_A.stl。现有六类基元：box、sphere、cylinder、cone/frustum、torus、capsule。产物为 JSON CSG 和原坐标系 STL。')

s=Slide(2,'相关工作','几何表示、代码生成与经验学习，是三条互补路线','本项目探索 RSI 在几何重建中的适配；尚未与这些方法进行同条件定量比较')
rows=[('UCSG-Net [1]','CSG 表示','学习基元参数和布尔树结构；不依赖完整树监督。'),('CAD-Recode [2]','CAD 代码生成','点云 → Python 草图 / 拉伸程序；使用训练后的模型。'),('StepCAD [3]','几何引导搜索','目标及中间几何条件下生成动作，用 IoU 引导程序搜索。'),('RSIAgent [4]','经验与记忆','课程、执行和验证形成闭环；冻结记忆后复用，不更新权重。')]
for i,(name,tag,desc) in enumerate(rows):
 y=191+i*88;s.rect(45,y,1190,76,WHITE,radius=True);s.text(63,y+19,256,40,name,25,BLUE,True);s.text(327,y+23,174,30,tag,19,TEAL);s.text(509,y+16,708,61,desc,21)
s.banner(560,'本项目：复用 RSI 调度框架，新增 CSG 工具、合成几何课程和独立数值验证')
for i,key in enumerate(refs):
 title,authors,url,doi=refs[key];x=46+(i%2)*600;y=620+(i//2)*25
 s.text(x,y,585,25,f'[{key}] {authors.split(",")[0]} · arXiv:{url.rsplit("/",1)[1]}',15,GRAY,link=url)
s.note('\n\n'.join(f'[{k}] {v[0]}. {v[1]}. {v[2]} DOI: {v[3]}' for k,v in refs.items())+'\n核对日期：2026-10-10。来源为官方 arXiv 摘要及 RSIAgent 上游 README。这里只比较研究路线，不搬用论文分数作为本项目结果；StepCAD 为预印本。')

s=Slide(3,'方法流程','先广泛练习，再针对目标学习，最后冻结记忆重建','Curriculum 设计练习 · Actor 写程序并执行 · Verifier 独立检查结果 [4]')
for x,num,title,desc in [(45,'01','BRS：广泛探索','合成多样几何练习\n最多 8 项，并行 2 项\n同一初始记忆 → 独立验证\n原 Actor 按顺序整理经验'),(455,'02','DRS：目标驱动','重建当前目标 → 验证\n若有缺陷，补充针对性练习\n更新记忆后由新 Actor 再试\n当前最多 4 个目标周期'),(865,'03','冻结记忆与评估','新 Actor / 新执行环境\n记忆只读，最多 3 次修订\n先完成工作阶段验证\n再运行独立最终评分')]:
 s.rect(x,202,370,255,WHITE,radius=True);s.text(x+20,219,60,45,num,32,TEAL);s.text(x+20,273,338,43,title,25,BLUE,True);s.text(x+20,324,338,129,desc,21)
s.text(421,310,32,45,'→',28,TEAL);s.text(831,310,32,45,'→',28,TEAL)
s.rect(45,481,1190,113,LIGHT,radius=True)
s.text(65,498,550,88,'记忆：文本笔记与脚本，不更新模型权重\n每个 case 独立学习，不共享前例记忆',22)
s.text(652,498,564,88,'执行：Docker 中分析 / 优化 / 导出 mesh\n模型：azure/gpt-5.6-sol，角色上下文隔离',21)
s.banner(615,'当前验证：64³ IoU ≥ 0.95 + 表面 P95 ≤ 对角线的 1% + 网格闭合')
s.note('实现证据：rsi_mesh/pipeline.py、runtime.py 和 configs/rsi_standard.yaml。调用上游 evolve_parallel_phase1、run_self_evolving_loop、core.loop.run_attempt。课程 Actor 只见合成 mesh，不见生成树。Verifier 不读取 Actor 记忆，但各角色使用同一模型，因此不是跨模型独立性。最终评分为96³，保留48³历史口径。所有案例学习预算2小时、最终30分钟。来源 [4] '+refs['4'][2])

s=Slide(4,'实验结果 · 九例','9 个样例中，7 个完成全流程并通过最终几何门槛','Fusion360 Gallery 项目内选取样例：easy / medium / hard 各 3 个，单次运行')
for x,val,label in [(45,'7 / 9','全流程完成并达标'),(451,'3 / 3','hard 样例达标'),(856,f'${cc:.2f}','九例模型费用估算')]:
 s.rect(x,189,379,86,WHITE,radius=True);s.text(x+17,197,140,58,val,36,TEAL,True);s.text(x+161,215,208,35,label,19)
s.text(48,299,550,34,'最终 IoU（96³）与基元数',23,BLUE,True)
xs=[60,300,440];s.rect(45,341,535,32,LIGHT);s.text(xs[0],344,225,29,'样例',17);s.text(xs[1],344,130,29,'IoU',17);s.text(xs[2],344,120,29,'基元数',17)
completed=[(n,v) for n,v in results if v['status']=='completed']
for i,(name,v) in enumerate(completed):
 y=378+i*32;m=v['metrics'];short='_'.join(name.split('_')[:2]);s.text(xs[0],y,234,30,short,18);s.text(xs[1],y,135,30,f'{m["voxel_iou_96"]:.4f}',18);s.text(xs[2],y,110,30,str(m['primitive_count']),18)
s.rect(620,296,615,309,WHITE,radius=True);s.text(639,309,570,34,'hard_109857：22 个基元，IoU 0.9839',22,BLUE)
s.image(assets/'hard_input.png',638,353,280,225);s.image(assets/'hard_output.png',941,353,280,225);s.text(720,570,160,30,'输入 mesh',17,GRAY);s.text(1010,570,190,30,'重建 mesh',17,GRAY)
s.text(47,616,1180,57,'两例中断：easy_129012 网格导出不闭合；easy_23853 最终阶段遇到服务 502/503。\n结果显示流程可行；尚无同预算无 RSI 对照，不能据此认定 RSI 带来提升。',19,ORANGE)
s.note('数据来源：rsi_runs/standard_9_20261009_trial01/cases/*/result.json 和 usage.jsonl。7个completed案例的96³ IoU均>=0.95且表面P95/对角线<=0.01。两例无最终分数，不能从分母剔除。选择性展示hard_109857是案例可视化，不是最佳整体质量证明。九例共11,756,251 tokens，官方标准价估算32.520142美元，含中断案例的已记录用量。')

s=Slide(5,'实验结果 · Voron','提高上限，未充分转化为更细的几何表达','同一 STL；两轮独立学习。两轮 64³ 验证均 PASS，但最终 96³ IoU 均 < 0.95')
for x,img,label in [(45,'voron_input','输入 mesh'),(450,'voron_24','上限 24：实际 20 个基元'),(855,'voron_48','上限 48：实际 22 个基元')]:
 s.rect(x,186,380,257,WHITE,radius=True);s.text(x+15,196,350,35,label,20,BLUE);s.image(assets/f'{img}.png',x+15,235,350,194)
headers=['指标','上限 24','上限 48']; xs=[59,260,421];s.rect(45,458,545,33,LIGHT)
for x,t in zip(xs,headers):s.text(x,461,175,28,t,17)
vals=[('最终 IoU 96³',f'{a["metrics"]["voxel_iou_96"]:.4f}',f'{b["metrics"]["voxel_iou_96"]:.4f}'),('P95 / 对角线','0.488%','0.549%'),('耗时（分钟）','84.8','76.0'),('估算费用',f'${ca:.2f}',f'${cb:.2f}')]
for i,row in enumerate(vals):
 for x,t in zip(xs,row):s.text(x,501+i*34,180,32,t,19)
s.text(625,465,592,185,'IoU 仅增加 0.0063，P95 略变差\n可见缺陷：漏孔、圆角与连接处近似\n实际 22 个基元仍低于原上限 24\n不能把两轮变化直接归因于容量增加',22,ORANGE)
s.note('数据来源：voron_motor_mount_trial01 与 voron_motor_mount_48_trial01 的 final/metrics.json、result.json 和 usage.jsonl。精确96³ IoU分别0.9367009335、0.9430131174，差0.0063121839。上限48实验树深也从12增加到16；课程和采样执行轨迹不同，无因果消融结论。单一斜视仅辅助展示，完整多视角比较见原报告。第一轮4,175,668 tokens，估算6.688496美元；第二轮3,470,814 tokens，估算6.046386美元。')

s=Slide(6,'未来计划','下一步：让验证反馈真正约束功能性细节','以下为待实施方案，当前已暂停新实验；不将计划写成已取得的结果')
for x,y,num,title,desc in [(45,192,'01','统一验证与最终口径','修复写死的 64³ 验证\n采用多分辨率 / 偏移采样\n分别记录工作 PASS 与最终达标'),(655,192,'02','从整体误差到局部缺陷','检测孔洞、连通性和截面差异\n定位缺失 / 多余材料\n用实体图与剖面辅助反馈'),(45,382,'03','增强表达与针对性课程','探索圆角基元与平滑连接\n围绕未通过的细节补学\n保留最佳候选，避免修改退化'),(655,382,'04','验证 RSI 的实际收益','同表示 / 同预算：有无 RSI 对照\nBRS、DRS、记忆消融与重复运行\n同时报告精度、成功率和 cost')]:
 s.rect(x,y,580,172,WHITE,radius=True);s.text(x+17,y+13,52,41,num,29,TEAL);s.text(x+80,y+17,480,40,title,24,BLUE,True);s.text(x+21,y+66,544,99,desc,21)
s.banner(577,f'已完成三次实验：19.40M token，模型费用合计估算 ${cc+ca+cb:.2f}；后续运行另行启动')
s.text(46,637,1180,35,'费用为公开标准价估算，非网关账单；不含机器费用及未返回用量的失败请求。[价格来源见备注]',16,GRAY)
s.note('未来计划源自本次项目讨论，用户随后明确暂停新实验。尚未实现圆角基元、新细节检查和预算扩展。收费来源：https://developers.openai.com/api/docs/models/gpt-5.6-sol ，查询2026-10-10。每百万token普通输入4美元、缓存读0.40美元、缓存写5美元、输出20美元；推理已计入输出。三次已完成实验总计19,402,733 tokens，估算45.2550228美元。此处不包含预检、合成验收及开发调用。PPT由academic-research-suite的证据核查与可视化原则整理；未运行跨模型审稿。')

if overflow: raise RuntimeError('Text overflow: '+repr(overflow))
for s in slides:s.im.save(ROOT/f'slide_{s.num:02d}.png')
prs.save(ROOT/'RSI_Mesh2CAD_汇报.pptx')
slides[0].im.save(ROOT/'RSI_Mesh2CAD_汇报.pdf',save_all=True,append_images=[s.im for s in slides[1:]],resolution=192)
thumb=Image.new('RGB',(1280,1080),LIGHT)
for i,s in enumerate(slides):thumb.paste(s.im.resize((640,360)),((i%2)*640,(i//2)*360))
thumb.save(ROOT/'总览.png')
manifest={'slides':6,'language':'zh-CN','data_runs':[str(batch),str(old),str(new)],'references':refs,'total_tokens':tc+ta+tb,'estimated_usd':cc+ca+cb,'verification_date':'2026-10-10','preview':'PNG/PDF rendered from shared layout; PPTX uses editable text and shapes.'}
(ROOT/'sources.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
print('Generated 6 slides, PPTX, PDF, PNG previews; no text overflow.')
